from __future__ import annotations

import io

import pytest
from docx import Document
from pypdf import PdfWriter

from app.parsing.chunking import chunk_text, estimate_tokens
from app.parsing.extract import ParseError, extract_document, parse_curated_qa, validate_upload


def test_txt_locations_and_utf8_bom_are_preserved():
    result = extract_document("notes.txt", "\ufeffHeading\nUseful source text".encode(), 10, 1000)
    assert result.text.startswith("Heading")
    assert result.locations[0][2].line_start == 1


def test_docx_keeps_heading_as_citation_section():
    file = io.BytesIO()
    document = Document()
    document.add_heading("Retrieval", 1)
    document.add_paragraph("Retrieve approved passages.")
    document.save(file)
    result = extract_document("notes.docx", file.getvalue(), 10, 1000)
    assert "Retrieval" in result.text
    assert any(location.section == "Retrieval" for _, _, location in result.locations)


def test_scanned_pdf_is_classified_for_ocr():
    output = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.write(output)
    with pytest.raises(ParseError, match="needs OCR"):
        extract_document("scan.pdf", output.getvalue(), 10, 1000)


def test_chunking_respects_limit_and_preserves_section_location():
    text = "one two three four five\n\nsix seven eight nine ten\n\neleven twelve thirteen fourteen"
    chunks = chunk_text(
        text,
        [
            (
                0,
                len(text),
                type(
                    "Loc",
                    (),
                    {"page": None, "section": "Chunking", "line_start": None, "line_end": None},
                )(),
            )
        ],
        8,
        2,
    )
    assert chunks
    assert all(estimate_tokens(chunk.text) <= 8 for chunk in chunks)
    assert chunks[0].location.section == "Chunking"


def test_curated_qa_requires_complete_pairs():
    assert parse_curated_qa("Question: What is RAG?\nAnswer: Retrieval augmented generation.")
    with pytest.raises(ParseError, match="both a Q"):
        parse_curated_qa("Q: unfinished")


def test_file_signature_validation_rejects_mislabelled_file():
    with pytest.raises(ParseError, match="does not match"):
        validate_upload("bad.pdf", b"not a pdf", 100)
