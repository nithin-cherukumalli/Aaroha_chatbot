from __future__ import annotations

import io
import re
import zipfile

from docx import Document as DocxDocument
from pypdf import PdfReader

from app.domain import ExtractedDocument, SourceLocation


class ParseError(ValueError):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


def validate_upload(filename: str, data: bytes, max_bytes: int) -> str:
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension not in {"txt", "pdf", "docx"}:
        raise ParseError("unsupported_media_type", "Only TXT, PDF, and DOCX files are supported.")
    if not data or len(data) > max_bytes:
        raise ParseError("file_too_large", "The file is empty or exceeds the upload size limit.")
    signatures = {"pdf": data.startswith(b"%PDF-"), "docx": data.startswith(b"PK")}
    if extension in signatures and not signatures[extension]:
        raise ParseError(
            "invalid_file_signature", "The uploaded file does not match its extension."
        )
    return extension


def _text_location(text: str) -> list[tuple[int, int, SourceLocation]]:
    locations, position = [], 0
    for line_number, line in enumerate(text.splitlines(keepends=True), start=1):
        locations.append(
            (
                position,
                position + len(line),
                SourceLocation(line_start=line_number, line_end=line_number),
            )
        )
        position += len(line)
    return locations


def _extract_txt(data: bytes) -> ExtractedDocument:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ParseError("invalid_text_encoding", "TXT files must use UTF-8 encoding.") from exc
    if not text.strip():
        raise ParseError("empty_document", "No readable text was found in this file.")
    return ExtractedDocument(text=text, warnings=[], locations=_text_location(text))


def _extract_pdf(data: bytes, max_pages: int) -> ExtractedDocument:
    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise ParseError("invalid_pdf", "The PDF could not be opened.") from exc
    if reader.is_encrypted:
        raise ParseError("encrypted_pdf", "Encrypted PDFs cannot be processed.")
    if len(reader.pages) > max_pages:
        raise ParseError("pdf_page_limit", "The PDF exceeds the configured page limit.")
    parts, locations, warnings, start = [], [], [], 0
    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if not text.strip():
            warnings.append(f"Page {page_number} has no extractable text.")
            continue
        parts.append(text)
        locations.append((start, start + len(text), SourceLocation(page=page_number)))
        start += len(text) + 2
    text = "\n\n".join(parts)
    if not text.strip():
        raise ParseError("ocr_required", "This PDF appears scanned or image-only and needs OCR.")
    warnings.append(
        "PDF layout, figures, and columns may not preserve their original visual order."
    )
    return ExtractedDocument(text=text, warnings=warnings, locations=locations)


def _extract_docx(data: bytes) -> ExtractedDocument:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > 10_000 or sum(item.file_size for item in infos) > 100 * 1024 * 1024:
                raise ParseError(
                    "docx_zip_limit", "The DOCX archive is too large to process safely."
                )
            if any(
                item.filename.endswith(".rels") and b'TargetMode="External"' in archive.read(item)
                for item in infos
                if item.filename.endswith(".rels")
            ):
                raise ParseError(
                    "docx_external_relationship",
                    "DOCX files with external relationships are not accepted.",
                )
        document = DocxDocument(io.BytesIO(data))
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError("invalid_docx", "The DOCX file could not be opened.") from exc
    parts: list[str] = []
    locations: list[tuple[int, int, SourceLocation]] = []
    start, current_heading, ordinal = 0, None, 0
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        ordinal += 1
        if paragraph.style and paragraph.style.name.lower().startswith("heading"):
            current_heading = text
        parts.append(text)
        locations.append(
            (
                start,
                start + len(text),
                SourceLocation(section=current_heading or f"paragraph {ordinal}"),
            )
        )
        start += len(text) + 2
    for table_number, table in enumerate(document.tables, start=1):
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        text = "\n".join(row for row in rows if row.strip())
        if text:
            parts.append(text)
            locations.append(
                (
                    start,
                    start + len(text),
                    SourceLocation(
                        section=f"{current_heading or 'document'}, table {table_number}"
                    ),
                )
            )
            start += len(text) + 2
    text = "\n\n".join(parts)
    if not text.strip():
        raise ParseError("empty_document", "No readable text was found in this DOCX file.")
    return ExtractedDocument(text=text, warnings=[], locations=locations)


def extract_document(
    filename: str, data: bytes, max_pages: int, max_characters: int
) -> ExtractedDocument:
    extension = filename.rsplit(".", 1)[-1].lower()
    result = {
        "txt": _extract_txt,
        "pdf": lambda b: _extract_pdf(b, max_pages),
        "docx": _extract_docx,
    }[extension](data)
    if len(result.text) > max_characters:
        raise ParseError("extracted_text_limit", "The extracted text exceeds the configured limit.")
    return result


QA_PATTERN = re.compile(
    r"(?ims)^\s*(?:Q|Question)\s*:\s*(.+?)\s*^\s*(?:A|Answer)\s*:\s*(.+?)(?=^\s*(?:Q|Question)\s*:|\Z)"
)


def parse_curated_qa(text: str) -> list[tuple[str, str]]:
    pairs = [(question.strip(), answer.strip()) for question, answer in QA_PATTERN.findall(text)]
    labels = re.findall(r"(?im)^\s*(?:Q|Question|A|Answer)\s*:", text)
    if labels and (not pairs or len(labels) != len(pairs) * 2):
        raise ParseError(
            "incomplete_qa", "Each curated question must contain both a Q: and A: block."
        )
    if not pairs:
        raise ParseError(
            "missing_qa", "Curated Q&A files must use Q:/A: or Question:/Answer: labels."
        )
    return pairs
